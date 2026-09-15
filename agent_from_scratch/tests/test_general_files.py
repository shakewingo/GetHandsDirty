import json
import os
from pathlib import Path
import stat
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from agent_from_scratch.agent import Agent
from agent_from_scratch.llm import ToolCall, LLM, LLMResponse, ResponseType
from agent_from_scratch.tools.files import EditFileTool, ListFilesTool, ReadFileTool, WriteFileTool
from agent_from_scratch.tools.register import ToolRegistry


class GeneralFileTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        self.workspace = self.root / "workspace"
        self.workspace.mkdir()
        self.read = ReadFileTool(self.workspace)
        self.write = WriteFileTool(self.workspace)
        self.edit = EditFileTool(self.workspace)
        self.list = ListFilesTool(self.workspace)

    def test_absolute_relative_home_and_symlink_paths(self):
        target = self.root / "outside.txt"
        target.write_text("outside")
        (self.workspace / "alias").symlink_to(target)
        for path in (str(target), "../outside.txt", "alias"):
            self.assertEqual(self.read.invoke({"path": path}).output["content"], "1| outside")
        with patch.dict(os.environ, {"HOME": str(self.root)}):
            self.assertEqual(self.read.invoke({"path": "~/outside.txt"}).output["content"], "1| outside")
        self.assertTrue(self.edit.invoke({"path": "alias", "old_text": "outside", "new_text": "updated"}).ok)
        self.assertTrue((self.workspace / "alias").is_symlink())
        self.assertEqual(target.read_text(), "updated")

    def test_optional_confinement_blocks_reads_and_mutations(self):
        target = self.root / "outside.txt"
        target.write_text("keep")
        (self.workspace / "alias").symlink_to(target)
        for cls, extra in [(ReadFileTool, {}), (WriteFileTool, {"content": "bad"}),
                           (EditFileTool, {"old_text": "keep", "new_text": "bad"}), (ListFilesTool, {})]:
            tool = cls(self.workspace, restrict_to_workspace=True)
            for path in ("../outside.txt", "alias"):
                self.assertEqual(tool.invoke({"path": path, **extra}).error_code, "denied")
        self.assertEqual(target.read_text(), "keep")

    def test_line_ranges_empty_files_and_character_continuation(self):
        target = self.workspace / "notes.txt"
        target.write_text("first\n你好🙂" * 70 + "\nlast\n")
        all_lines = target.read_text().splitlines()
        reader = ReadFileTool(self.workspace, max_chars=64)
        recovered = {}
        offset, column = 1, 0
        for _ in range(100):
            result = reader.invoke({"path": "notes.txt", "offset": offset, "column": column, "limit": 3})
            self.assertTrue(result.ok, result.error_message)
            self.assertLessEqual(len(result.output["content"]), 64)
            for line in result.output["content"].splitlines():
                number, text = line.split("| ", 1)
                recovered[int(number)] = recovered.get(int(number), "") + text
            if result.output["eof"]:
                break
            cursor = result.output["next_offset"], result.output["next_column"]
            self.assertGreater(cursor, (offset, column))
            offset, column = cursor
        self.assertEqual([recovered[i] for i in sorted(recovered)], all_lines)
        target.write_text("")
        self.assertEqual(self.read.invoke({"path": "notes.txt"}).output["content"], "")
        self.assertEqual(self.read.invoke({"path": "notes.txt", "offset": 2}).error_code, "invalid_arguments")

    def test_one_long_line_is_not_lost_at_the_output_cap(self):
        text = "你好🙂" * 100
        (self.workspace / "long.txt").write_text(text)
        reader = ReadFileTool(self.workspace, max_chars=32)
        parts, column = [], 0
        while True:
            output = reader.invoke({"path": "long.txt", "column": column}).output
            parts.append(output["content"].split("| ", 1)[1])
            if output["eof"]:
                break
            self.assertEqual(output["next_offset"], 1)
            self.assertGreater(output["next_column"], column)
            column = output["next_column"]
        self.assertEqual("".join(parts), text)

    def test_encoding_and_invalid_ranges_or_binary_inputs(self):
        (self.workspace / "utf16.txt").write_bytes("你好\nworld".encode("utf-16"))
        result = self.read.invoke({"path": "utf16.txt", "encoding": "utf-16", "offset": 2, "limit": 1})
        self.assertEqual(result.output["content"], "2| world")
        os.mkfifo(self.workspace / "pipe")
        (self.workspace / "binary").write_bytes(b"\0x")
        for path in ("", "missing", ".", "pipe", "binary"):
            self.assertFalse(self.read.invoke({"path": path}).ok)
        for args in ({"offset": 0}, {"offset": True}, {"limit": 0}, {"chunk_size": 4}, {"pages": "1"}):
            self.assertEqual(self.read.invoke({"path": "utf16.txt", **args}).error_code, "invalid_arguments")

    def test_large_writes_append_noop_and_permissions(self):
        target = self.workspace / "nested/program.py"
        text = "# source\n" * 3000
        self.assertTrue(self.write.invoke({"path": "nested/program.py", "content": text}).ok)
        target.chmod(0o755)
        version = self.read.invoke({"path": "nested/program.py"}).output["version"]
        result = self.write.invoke({"path": "nested/program.py", "content": text})
        self.assertFalse(result.output["changed"])
        self.assertEqual(result.output["version"], version)
        result = self.write.invoke({"path": "nested/program.py", "content": "print('hi')\n", "append": True})
        self.assertTrue(result.ok, result.error_message)
        self.assertEqual(target.read_text(), text + "print('hi')\n")
        self.assertEqual(stat.S_IMODE(target.stat().st_mode), 0o755)

    def test_stale_version_does_not_overwrite_an_external_change(self):
        target = self.workspace / "config.txt"
        target.write_text("old")
        version = self.read.invoke({"path": "config.txt"}).output["version"]
        target.write_text("external")
        for tool, args in [(self.write, {"content": "bad"}),
                           (self.edit, {"old_text": "external", "new_text": "bad"})]:
            result = tool.invoke({"path": "config.txt", "expected_version": version, **args})
            self.assertFalse(result.ok)
            self.assertIn("version", result.error_message)
        self.assertEqual(target.read_text(), "external")

    def test_replacement_failure_preserves_file_and_cleans_temp(self):
        target = self.workspace / "config.txt"
        target.write_text("old")
        with patch.object(Path, "replace", side_effect=OSError("disk failure")):
            self.assertFalse(self.edit.invoke({"path": "config.txt", "old_text": "old", "new_text": "new"}).ok)
        self.assertEqual(target.read_text(), "old")
        self.assertEqual(list(self.workspace.iterdir()), [target])

    def test_change_during_temporary_write_preserves_the_external_file(self):
        target = self.workspace / "config.txt"
        target.write_text("old")
        with patch("agent_from_scratch.tools.files.os.fsync", side_effect=lambda fd: target.write_text("external")):
            result = self.write.invoke({"path": "config.txt", "content": "new"})
        self.assertFalse(result.ok)
        self.assertIn("changed before replacement", result.error_message)
        self.assertEqual(target.read_text(), "external")
        self.assertEqual(list(self.workspace.iterdir()), [target])

    def test_null_optional_edit_fields_are_accepted_but_wrong_line_is_not(self):
        target = self.workspace / "config.py"
        target.write_text("# comment\nretries = 2\n")
        args = {"path": "config.py", "old_text": "retries = 2", "new_text": "retries = 3",
                "occurrence": None, "line_hint": None, "expected_replacements": None, "expected_version": None}
        self.assertFalse(self.edit.invoke({**args, "line_hint": 1}).ok)
        result = self.edit.invoke(args)
        self.assertTrue(result.ok, result.error_message)
        self.assertEqual(target.read_text(), "# comment\nretries = 3\n")

    def test_file_size_cap_applies_before_new_directories_or_mutations(self):
        writer = WriteFileTool(self.workspace, max_file_bytes=4)
        self.assertFalse(writer.invoke({"path": "new/a", "content": "too long"}).ok)
        self.assertFalse((self.workspace / "new").exists())
        (self.workspace / "big").write_text("too long")
        self.assertFalse(ReadFileTool(self.workspace, max_file_bytes=4).invoke({"path": "big"}).ok)

    def test_mid_read_modification_is_not_returned_as_success(self):
        target = self.workspace / "a"
        target.write_text("before")
        before = target.stat()
        changed = SimpleNamespace(**{key: getattr(before, key) for key in (
            "st_dev", "st_ino", "st_size", "st_mtime_ns", "st_ctime_ns")})
        changed.st_mtime_ns += 1
        with patch("agent_from_scratch.tools.files.os.fstat", side_effect=[before, changed]):
            self.assertFalse(self.read.invoke({"path": "a"}).ok)

    def test_edit_selectors_reject_ambiguity_and_count_mismatch(self):
        target = self.workspace / "config.txt"
        original = "x=1\ny=2\nx=1\n"
        args = {"path": "config.txt", "old_text": "x=1", "new_text": "x=3"}
        for extra in ({}, {"occurrence": 3}, {"line_hint": 2},
                      {"replace_all": True, "expected_replacements": 3}, {"replace_all": True, "occurrence": 1}):
            target.write_text(original)
            self.assertFalse(self.edit.invoke({**args, **extra}).ok)
            self.assertEqual(target.read_text(), original)
        for extra in ({"occurrence": 2}, {"line_hint": 3}):
            target.write_text(original)
            result = self.edit.invoke({**args, **extra, "expected_replacements": 1})
            self.assertTrue(result.ok, result.error_message)
            self.assertEqual(target.read_text(), "x=1\ny=2\nx=3\n")
        target.write_text(original)
        result = self.edit.invoke({**args, "replace_all": True, "expected_replacements": 2})
        self.assertEqual(result.output["replacements"], 2)
        self.assertEqual(target.read_text(), "x=3\ny=2\nx=3\n")

    def test_edit_create_delete_crlf_and_noop(self):
        args = {"path": "new.txt", "old_text": "", "new_text": "a\r\nb\r\nlast\r\n"}
        self.assertTrue(self.edit.invoke(args).output["created"])
        self.assertFalse(self.edit.invoke(args).ok)
        target = self.workspace / "new.txt"
        result = self.edit.invoke({"path": "new.txt", "old_text": "a\nb", "new_text": "x\ny"})
        self.assertTrue(result.ok, result.error_message)
        self.assertEqual(target.read_bytes(), b"x\r\ny\r\nlast\r\n")
        result = self.edit.invoke({"path": "new.txt", "old_text": "last\n", "new_text": ""})
        self.assertTrue(result.ok)
        self.assertEqual(target.read_bytes(), b"x\r\ny\r\n")
        self.assertFalse(self.edit.invoke({"path": "new.txt", "old_text": "x", "new_text": "x"}).output["changed"])

    def test_listing_recursion_pagination_ignored_dirs_and_symlink_cycle(self):
        (self.workspace / "a").mkdir()
        (self.workspace / "a/source.py").write_text("123")
        (self.workspace / ".git").mkdir()
        (self.workspace / ".git/config").write_text("hidden")
        (self.workspace / "z").symlink_to(self.workspace, target_is_directory=True)
        entries, offset = [], 0
        while True:
            result = self.list.invoke({"path": ".", "recursive": True, "max_entries": 1, "offset": offset})
            self.assertTrue(result.ok, result.error_message)
            entries.extend(result.output["entries"])
            offset = result.output["next_offset"]
            if offset is None:
                break
        self.assertEqual([entry["name"] for entry in entries], ["a", "a/source.py", "z"])
        self.assertEqual([entry["type"] for entry in entries], ["directory", "file", "symlink"])
        self.assertEqual(entries[1]["size_bytes"], 3)
        all_entries = self.list.invoke({"path": ".", "recursive": True, "include_ignored": True}).output["entries"]
        self.assertIn(".git/config", [entry["name"] for entry in all_entries])

    def test_read_edit_read_uses_the_ordinary_loop_and_traces_versions(self):
        (self.workspace / "config.txt").write_text("value=1\n")
        model = Mock(spec=LLM)
        model.settings.return_value = {}
        model.generate.side_effect = [
            LLMResponse('assistant', '', ResponseType.tool_call, tool_calls=[ToolCall('read_file', {'path': 'config.txt'})]),
            LLMResponse('assistant', '', ResponseType.tool_call, tool_calls=[ToolCall('edit_file', {'path': 'config.txt', 'old_text': 'value=1', 'new_text': 'value=2'})]),
            LLMResponse('assistant', '', ResponseType.tool_call, tool_calls=[ToolCall('read_file', {'path': 'config.txt'})]),
            LLMResponse("assistant", "Updated", ResponseType.direct),
        ]
        result = Agent(model, registry=ToolRegistry([self.read, self.edit])).run_turn("Update the value")
        observations = [json.loads(m["content"]) for m in result.messages if m["role"] == "tool"]
        self.assertEqual([o["ok"] for o in observations], [True, True, True])
        self.assertNotEqual(observations[0]["output"]["version"], observations[2]["output"]["version"])
        self.assertEqual(observations[2]["output"]["content"], "1| value=2")


class DocumentReadTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.read = ReadFileTool(self.root)

    def test_docx_paragraphs_and_tables_preserve_order_and_can_be_paged(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Start here")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "command"
        table.cell(0, 1).text = "python app.py"
        document.add_paragraph("End here")
        document.save(self.root / "guide.docx")
        first = self.read.invoke({"path": "guide.docx", "limit": 2}).output
        self.assertEqual(first["content"], "1| Start here\n2| command | python app.py")
        self.assertFalse(first["eof"])
        last = self.read.invoke({"path": "guide.docx", "offset": first["next_offset"]}).output
        self.assertEqual(last["content"], "3| End here")
        self.assertTrue(last["eof"])

    def test_xlsx_retains_sheet_names_row_numbers_and_formulas(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["name", "value"])
        sheet.append(["alpha", 3])
        sheet.append(["total", "=SUM(B2:B2)"])
        workbook.save(self.root / "data.xlsx")
        workbook.close()
        result = self.read.invoke({"path": "data.xlsx"})
        self.assertTrue(result.ok, result.error_message)
        self.assertIn("Sheet: Data", result.output["content"])
        self.assertIn("Row 2: alpha\t3", result.output["content"])
        self.assertIn("=SUM(B2:B2)", result.output["content"])

    def test_pptx_retains_slide_text_and_table_cells(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Tool design"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "loop", "trace"
        presentation.save(self.root / "slides.pptx")
        result = self.read.invoke({"path": "slides.pptx"})
        self.assertTrue(result.ok, result.error_message)
        self.assertIn("Slide 1", result.output["content"])
        self.assertIn("Tool design", result.output["content"])
        self.assertIn("loop | trace", result.output["content"])

    def test_pdf_text_pages_and_invalid_ranges(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        for number in range(1, 4):
            page = writer.add_blank_page(width=300, height=200)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            contents = DecodedStreamObject()
            contents.set_data(f"BT /F1 12 Tf 30 100 Td (Page {number} evidence) Tj ET".encode())
            page[NameObject("/Contents")] = contents
        writer.write(self.root / "pages.pdf")
        writer.close()
        result = self.read.invoke({"path": "pages.pdf", "pages": "2"})
        self.assertTrue(result.ok, result.error_message)
        self.assertIn("Page 2 evidence", result.output["content"])
        self.assertNotIn("Page 1 evidence", result.output["content"])
        self.assertEqual(result.output["next_pages"], "3-3")
        for pages in ("0", "4", "2-1", "1,3", "x"):
            self.assertEqual(self.read.invoke({"path": "pages.pdf", "pages": pages}).error_code, "invalid_arguments")

    def test_pdf_default_page_cap_empty_and_encrypted_documents(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        for _ in range(21):
            writer.add_blank_page(width=100, height=100)
        writer.write(self.root / "pages.pdf")
        result = self.read.invoke({"path": "pages.pdf"})
        self.assertTrue(result.ok)
        self.assertEqual(result.output["next_pages"], "21-21")
        self.assertFalse(result.output["document_eof"])
        self.assertIn("No extractable text", result.output["content"])
        self.assertEqual(self.read.invoke({"path": "pages.pdf", "pages": "1-21"}).error_code, "invalid_arguments")
        writer.encrypt("test-password")
        writer.write(self.root / "locked.pdf")
        writer.close()
        self.assertFalse(self.read.invoke({"path": "locked.pdf"}).ok)

    def test_images_report_metadata_without_claiming_visual_access(self):
        from PIL import Image

        Image.new("RGB", (12, 8)).save(self.root / "picture.png")
        result = self.read.invoke({"path": "picture.png"})
        self.assertTrue(result.ok)
        self.assertEqual((result.output["width"], result.output["height"]), (12, 8))
        self.assertFalse(result.output["visual_content_available"])
        self.assertNotIn("content", result.output)

    def test_corrupt_document_is_an_error_and_next_read_still_works(self):
        (self.root / "bad.docx").write_bytes(b"not a zip archive")
        self.assertFalse(self.read.invoke({"path": "bad.docx"}).ok)
        (self.root / "good.txt").write_text("good")
        self.assertTrue(self.read.invoke({"path": "good.txt"}).ok)


if __name__ == "__main__":
    unittest.main()
