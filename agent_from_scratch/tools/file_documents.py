"""Text extraction for read_file, following nanobot's document adapters.

Formats/layout/images are not reconstructed. Each adapter operates on the captured
file bytes, so an external edit cannot mix versions during document extraction.
"""

from collections.abc import Iterator
from io import BytesIO
import re
from zipfile import ZipFile

from .base import ToolErrorCode, ToolExecutionError


def image_metadata(raw: bytes) -> dict:
    from PIL import Image

    with Image.open(BytesIO(raw)) as image:
        return {"format": image.format, "width": image.width, "height": image.height,
                "mode": image.mode, "visual_content_available": False,
                "note": "Image metadata only. This text-only agent cannot inspect pixels or perform OCR."}


def document_lines(raw: bytes, suffix: str, pages: str | None) -> tuple[Iterator[str], dict]:
    metadata: dict = {"format": suffix.lstrip("."),
                "extraction_note": "Extracted text only; images, layout and OCR are not included."}
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted:
            raise ToolExecutionError(ToolErrorCode.EXECUTION_ERROR, "Encrypted PDFs require a decrypted copy.")
        total = len(reader.pages)
        start, end = 1, min(20, total)
        if pages is not None:
            match = re.fullmatch(r"\s*(\d+)(?:-(\d+))?\s*", pages)
            if not match:
                raise ToolExecutionError(ToolErrorCode.INVALID_ARGUMENTS, "Use pages='2' or pages='2-5'.")
            start, end = int(match[1]), int(match[2] or match[1])
            if not 1 <= start <= end <= total or end - start + 1 > 20:
                raise ToolExecutionError(ToolErrorCode.INVALID_ARGUMENTS, "Choose 1–20 existing PDF pages.")
        metadata.update(total_pages=total, pages=f"{start}-{end}" if total else None,
                        next_pages=f"{end + 1}-{min(end + 20, total)}" if end < total else None)

        def pdf_lines():
            for number in range(start, end + 1):
                page = reader.pages[number - 1]
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 10 * 1024 * 1024:
                    raise ToolExecutionError(ToolErrorCode.EXECUTION_ERROR, f"PDF page {number} exceeds the text-extraction limit.")
                yield f"--- Page {number} ---"
                text = page.extract_text() or ""
                yield from text.splitlines() or ["[No extractable text on this page; OCR is not available.]"]

        return pdf_lines(), metadata

    # Office files are ZIP containers. Check declared expanded size before parsing XML.
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 256 * 1024 * 1024:
            raise ToolExecutionError(ToolErrorCode.EXECUTION_ERROR, "Office archive exceeds extraction limits.")

    def office_lines():
        if suffix == ".docx":
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            document = Document(BytesIO(raw))
            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    yield from block.text.splitlines() or [""]
                elif isinstance(block, Table):
                    for row in block.rows:
                        yield " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            try:
                cells_seen = 0
                for sheet in workbook:
                    yield f"--- Sheet: {sheet.title} (formula text; not recalculated) ---"
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        cells_seen += len(row)
                        if cells_seen > 1000000:
                            raise ToolExecutionError(ToolErrorCode.EXECUTION_ERROR, "Spreadsheet scan exceeds one million cells.")
                        if any(value is not None for value in row):
                            yield f"Row {number}: " + "\t".join(str(value) if value is not None else "" for value in row)
            finally:
                workbook.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            def shapes_text(shapes):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        yield from shapes_text(shape.shapes)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            yield " | ".join(cell.text for cell in row.cells)
                    elif shape.has_text_frame:
                        yield from shape.text.splitlines()

            presentation = Presentation(BytesIO(raw))
            for number, slide in enumerate(presentation.slides, 1):
                yield f"--- Slide {number} ---"
                yield from shapes_text(slide.shapes)

    return office_lines(), metadata
